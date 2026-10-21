"""Extract page-tagged text from PDF, DOCX, and PPTX uploads.
Scanned PDFs fall back to multimodal OCR via OpenRouter."""

from __future__ import annotations

import base64
import io
import re
from collections import Counter

from fastapi import HTTPException

from app.services.pdf import (
    MIN_TEXT_CHARS,
    _join_hyphenated,
    _normalize_whitespace,
)

PDF_MAGIC = b"%PDF"
DOCX_MAGIC = b"PK"  # zip
MAX_OCR_PAGES = 12


def _ext(filename: str) -> str:
    name = (filename or "").rsplit(".", 1)
    return name[-1].lower() if len(name) == 2 else ""


def detect_kind(filename: str, data: bytes) -> str:
    ext = _ext(filename)
    if data.startswith(PDF_MAGIC) or ext == "pdf":
        return "pdf"
    if ext in ("docx",) and data[:2] == DOCX_MAGIC:
        return "docx"
    if ext in ("pptx",) and data[:2] == DOCX_MAGIC:
        return "pptx"
    if ext in ("ppt",):
        raise HTTPException(
            status_code=400,
            detail="Legacy .ppt is not supported. Save as .pptx or PDF and upload again.",
        )
    raise HTTPException(
        status_code=400,
        detail="Supported uploads: PDF, DOCX, and PPTX.",
    )


def extract_docx_pages(data: bytes) -> list[dict]:
    try:
        from docx import Document
    except ImportError as exc:
        raise HTTPException(status_code=500, detail="DOCX support is not installed.") from exc
    try:
        doc = Document(io.BytesIO(data))
    except Exception as exc:
        raise HTTPException(status_code=400, detail="Could not read this Word document.") from exc
    paras = [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
            if cells:
                paras.append(" | ".join(cells))
    text = _normalize_whitespace("\n".join(paras))
    if len(text) < MIN_TEXT_CHARS:
        raise HTTPException(
            status_code=400,
            detail="This Word document has almost no readable text.",
        )
    return [{"page": 1, "text": text}]


def extract_pptx_pages(data: bytes) -> list[dict]:
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise HTTPException(status_code=500, detail="PPTX support is not installed.") from exc
    try:
        prs = Presentation(io.BytesIO(data))
    except Exception as exc:
        raise HTTPException(
            status_code=400, detail="Could not read this PowerPoint file."
        ) from exc
    pages: list[dict] = []
    for i, slide in enumerate(prs.slides, start=1):
        bits: list[str] = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text and shape.text.strip():
                bits.append(shape.text.strip())
        text = _normalize_whitespace("\n".join(bits))
        if text:
            pages.append({"page": i, "text": text})
    total = sum(len(p["text"]) for p in pages)
    if total < MIN_TEXT_CHARS:
        raise HTTPException(
            status_code=400,
            detail="This presentation has almost no readable text. Export slides as PDF and try again.",
        )
    return pages


def extract_pdf_text_pages(data: bytes) -> list[dict]:
    """Text-layer PDF extraction (no OCR). May return sparse pages for scans."""
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
        raw_pages = [page.extract_text() or "" for page in reader.pages]
    except Exception as exc:
        raise HTTPException(status_code=400, detail="Could not read this PDF file.") from exc

    n_pages = len(raw_pages)
    line_counts: Counter[str] = Counter()
    page_lines = []
    for text in raw_pages:
        lines = [line.strip() for line in text.split("\n") if line.strip()]
        page_lines.append(lines)
        for line in set(lines):
            line_counts[line] += 1

    threshold = n_pages / 2
    repeated = {line for line, count in line_counts.items() if n_pages > 1 and count > threshold}

    pages = []
    for i, lines in enumerate(page_lines):
        kept = [line for line in lines if line not in repeated]
        text = _join_hyphenated("\n".join(kept))
        text = _normalize_whitespace(text)
        pages.append({"page": i + 1, "text": text})
    return pages


def _render_pdf_page_pngs(data: bytes, *, max_pages: int = MAX_OCR_PAGES) -> list[bytes]:
    try:
        import pymupdf
    except ImportError as exc:
        raise HTTPException(status_code=500, detail="OCR support is not installed.") from exc
    try:
        doc = pymupdf.open(stream=data, filetype="pdf")
    except Exception as exc:
        raise HTTPException(status_code=400, detail="Could not open this PDF for OCR.") from exc
    images: list[bytes] = []
    try:
        for i, page in enumerate(doc):
            if i >= max_pages:
                break
            pix = page.get_pixmap(matrix=pymupdf.Matrix(1.5, 1.5), alpha=False)
            images.append(pix.tobytes("png"))
    finally:
        doc.close()
    if not images:
        raise HTTPException(status_code=400, detail="This PDF has no pages to read.")
    return images


async def ocr_pdf_pages(data: bytes) -> list[dict]:
    """Use the multimodal OCR model to read scanned PDF pages."""
    from app.services.llm import LLMError, OpenAIClient, get_llm_client

    client = get_llm_client()
    if not isinstance(client, OpenAIClient):
        raise HTTPException(
            status_code=400,
            detail={
                "detail": (
                    "This PDF looks like scanned images. OCR needs a live AI provider. "
                    "Paste the text instead, or configure OpenRouter."
                ),
                "code": "SCANNED_PDF",
            },
        )

    pngs = _render_pdf_page_pngs(data)
    pages: list[dict] = []
    # Process in small batches to stay within free-tier limits
    batch_size = 2
    for start in range(0, len(pngs), batch_size):
        batch = pngs[start : start + batch_size]
        b64s = [base64.b64encode(p).decode("ascii") for p in batch]
        page_nos = list(range(start + 1, start + 1 + len(batch)))
        prompt = (
            "You are an OCR engine for course material. Transcribe ALL readable text "
            f"from these slide/page image(s) (pages {page_nos[0]}–{page_nos[-1]}). "
            "Preserve headings and bullet structure. Return plain text only. "
            "Separate pages with a line that is exactly: ===PAGE n=== where n is the page number."
        )
        try:
            text = await client.vision_text(prompt, b64s, max_tokens=3500)
        except LLMError as exc:
            raise HTTPException(status_code=502, detail=str(exc)) from exc
        split = re.split(r"===PAGE\s*(\d+)===", text)
        if len(split) >= 3:
            # split = [preamble, num, body, num, body, ...]
            for j in range(1, len(split), 2):
                num = int(split[j])
                body = _normalize_whitespace(split[j + 1] if j + 1 < len(split) else "")
                if body:
                    pages.append({"page": num, "text": body})
        else:
            # Fallback: assign whole batch text to first page in batch
            body = _normalize_whitespace(text)
            if body:
                pages.append({"page": page_nos[0], "text": body})

    total = sum(len(p["text"]) for p in pages)
    if total < MIN_TEXT_CHARS:
        raise HTTPException(
            status_code=400,
            detail="OCR could not read enough text from this PDF. Try a clearer scan or paste notes.",
        )
    return pages


async def extract_pages(filename: str, data: bytes) -> tuple[str, list[dict]]:
    """Return (kind, pages) for an uploaded file. Runs OCR for scanned PDFs."""
    kind = detect_kind(filename, data)
    if kind == "docx":
        return kind, extract_docx_pages(data)
    if kind == "pptx":
        return kind, extract_pptx_pages(data)

    pages = extract_pdf_text_pages(data)
    total = sum(len(p["text"]) for p in pages)
    if total >= MIN_TEXT_CHARS:
        return "pdf", pages
    # Scanned / image PDF → OCR
    ocr_pages = await ocr_pdf_pages(data)
    return "pdf", ocr_pages
